from __future__ import annotations

from pathlib import Path
from typing import Any

import markdown
from docx import Document
from docx.shared import Inches, Pt
from jinja2 import Environment, PackageLoader, StrictUndefined, select_autoescape
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from .selector import select_experiences, selected_skills


def _environment() -> Environment:
    return Environment(loader=PackageLoader("career_os", "templates"), autoescape=select_autoescape(default_for_string=False, default=False), undefined=StrictUndefined, trim_blocks=True, lstrip_blocks=True)


def _context(documents: dict[str, Any], profile_id: str | None) -> dict[str, Any]:
    experiences = select_experiences(documents, profile_id)
    return {"profile": documents.get("profile", {}).get("profile", {}), "target": documents.get("target_profiles", {}).get(profile_id, {}) if profile_id else {}, "experiences": experiences, "education": documents.get("education", {}).get("education", []), "certifications": documents.get("certifications", {}).get("certifications", []), "languages": documents.get("languages", {}).get("languages", []), "skills": selected_skills(documents, experiences, profile_id)}


def render_resume_markdown(documents: dict[str, Any], profile_id: str | None = None) -> str:
    return _environment().get_template("resume-ats-en.md.j2").render(**_context(documents, profile_id)).strip() + "\n"


def render_resume_html(documents: dict[str, Any], profile_id: str | None = None) -> str:
    body = markdown.markdown(render_resume_markdown(documents, profile_id), extensions=["sane_lists"])
    return _environment().get_template("resume-en.html.j2").render(body=body)


def _write_pdf(documents: dict[str, Any], output: Path, profile_id: str | None) -> None:
    ctx = _context(documents, profile_id); styles = getSampleStyleSheet()
    story = [Paragraph(ctx["profile"]["full_name"], styles["Title"])]
    if ctx["profile"].get("headline"): story += [Paragraph(ctx["profile"]["headline"], styles["Normal"]), Spacer(1, 4*mm)]
    story.append(Paragraph("Professional Experience", styles["Heading2"]))
    for item in ctx["experiences"]:
        story.append(Paragraph(f'{item["role"]} — {item["company"]}', styles["Heading3"]))
        end = "Present" if item.get("current") else item["dates"].get("end", "")
        story.append(Paragraph(f'{item["dates"]["start"]} – {end}', styles["Normal"]))
        bullets=[ListItem(Paragraph(v, styles["BodyText"])) for v in item.get("highlights", [])]
        if bullets: story.append(ListFlowable(bullets, bulletType="bullet", leftIndent=12))
    if ctx["skills"]: story += [Paragraph("Core Skills", styles["Heading2"]), Paragraph(", ".join(ctx["skills"]), styles["BodyText"])]
    if ctx["education"]:
        story.append(Paragraph("Education", styles["Heading2"]))
        for item in ctx["education"]: story.append(Paragraph(f'{item["degree"]} — {item["institution"]}', styles["BodyText"]))
    output.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(output), pagesize=A4, rightMargin=14*mm, leftMargin=14*mm, topMargin=12*mm, bottomMargin=12*mm).build(story)


def _write_docx(documents: dict[str, Any], output: Path, profile_id: str | None) -> None:
    ctx=_context(documents, profile_id); doc=Document(); section=doc.sections[0]
    section.top_margin=section.bottom_margin=Inches(0.55); section.left_margin=section.right_margin=Inches(0.65)
    doc.styles["Normal"].font.name="Arial"; doc.styles["Normal"].font.size=Pt(10)
    doc.add_heading(ctx["profile"]["full_name"],0)
    if ctx["profile"].get("headline"): doc.add_paragraph(ctx["profile"]["headline"])
    doc.add_heading("Professional Experience",1)
    for item in ctx["experiences"]:
        doc.add_heading(f'{item["role"]} — {item["company"]}',2)
        end="Present" if item.get("current") else item["dates"].get("end","")
        doc.add_paragraph(f'{item["dates"]["start"]} – {end}')
        for value in item.get("highlights",[]): doc.add_paragraph(value,style="List Bullet")
    if ctx["skills"]: doc.add_heading("Core Skills",1); doc.add_paragraph(", ".join(ctx["skills"]))
    if ctx["education"]:
        doc.add_heading("Education",1)
        for item in ctx["education"]: doc.add_paragraph(f'{item["degree"]} — {item["institution"]}')
    output.parent.mkdir(parents=True, exist_ok=True); doc.save(output)


def write_resume(documents: dict[str, Any], output: str | Path, profile_id: str | None = None, format: str | None = None) -> Path:
    output=Path(output); fmt=(format or output.suffix.lstrip(".") or "md").lower(); output.parent.mkdir(parents=True,exist_ok=True)
    if fmt in {"md","markdown"}: output.write_text(render_resume_markdown(documents,profile_id),encoding="utf-8")
    elif fmt=="html": output.write_text(render_resume_html(documents,profile_id),encoding="utf-8")
    elif fmt=="pdf": _write_pdf(documents,output,profile_id)
    elif fmt=="docx": _write_docx(documents,output,profile_id)
    else: raise ValueError(f"Unsupported format: {fmt}")
    return output


def write_resume_markdown(documents: dict[str, Any], output: str | Path) -> Path:
    return write_resume(documents,output,format="md")
